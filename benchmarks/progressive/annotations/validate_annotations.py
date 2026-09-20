"""Validate normalized owner-approved annotations without inventing verification."""
from pathlib import Path
import argparse
import hashlib
import json
import re
import openpyxl
from approval import content_digest

P = Path(__file__).resolve().parent
ROOT = P.parents[2]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--metadata-only', action='store_true', help='Use evidence manifest; do not require local media/source workbooks')
args = parser.parse_args()
a = [json.loads(line) for line in (P / 'ground_truth.jsonl').read_text().splitlines()]
inv = [json.loads(line) for line in (P / 'source_inventory.jsonl').read_text().splitlines()]
approval = json.loads((P / 'owner_approval.json').read_text())
manifest = json.loads((P / 'evidence_manifest.json').read_text())
files = {item['path']: item for item in manifest['files']}
assert len(files) == len(manifest['files'])
assert len({r['query_id'] for r in a}) == len(a)
assert len(inv) == 159
assert len({(r['source_records'][0]['workbook'], r['source_records'][0]['sheet'], r['source_records'][0]['row']) for r in inv}) == len(inv)
sha = {}
for r in inv:
    for source in r['source_records']:
        name, digest = source['workbook'], source['sha256']
        assert name not in sha or sha[name] == digest
        sha[name] = digest
if not args.metadata_only:
    for name, digest in sha.items():
        assert hashlib.sha256((ROOT / name).read_bytes()).hexdigest() == digest, name

checked = set()
def artifact(path, digest=None):
    assert path in files, path
    item = files[path]
    if digest:
        assert item['sha256'] == digest, path
    if not args.metadata_only and path not in checked:
        f = ROOT / path
        assert f.is_file() and f.stat().st_size == item['bytes'], path
        assert hashlib.sha256(f.read_bytes()).hexdigest() == item['sha256'], path
        checked.add(path)

for r in a:
    assert r['time_unit'] == 'ms'
    assert r['interval_convention'] == 'start_inclusive_end_exclusive'
    accepted = r['review_status'] == 'approved'
    if accepted:
        assert r['annotation_status'] == 'source_normalized'
        assert r['reviewer'] == 'dataset_owner'
        assert r['approval_id'] == approval['approval_id']
        assert content_digest(r) == r['approval_content_sha256'] == approval['record_digests'][r['query_id']]
        representative = r.get('duplicate_representative_query_id')
        eligible = r['task_type'] == 'TKIS' and (not representative or representative == r['query_id'])
        assert r['eligible_for_benchmark'] == eligible
        assert len(r['hints']) == 3
        assert any(t['target_intervals'] for t in r['targets'])
    else:
        assert r['annotation_status'] == 'unverified' and r['review_status'] == 'pending'
        assert not r['eligible_for_benchmark'] and not r['targets'] and not r['hints']
    ids = {e['evidence_id'] for e in r['evidence']}
    assert len(ids) == len(r['evidence'])
    for e in r['evidence']:
        artifact(e['artifact_path'], e['artifact_sha256'])
        artifact(e['clip_path'], e['clip_sha256'])
        assert e['continuous_playback_inspected'] is False and e['audio_inspected'] is False
        for sample in e['sampled_frames']:
            artifact(sample['path'])
    for k, hint in enumerate(r['hints']):
        assert hint['cumulative_text'] == ' '.join(h['delta_text'] for h in r['hints'][:k + 1])
        assert hint['evidence_ids'] and set(hint['evidence_ids']) <= ids
        assert hint['revealed_at_ms'] is None
        assert not re.search(r'L\d{2}_V\d+|\d{2}:\d{2}', hint['delta_text'])
        assert hint['review_status'] == r['review_status']
    for target in r['targets']:
        assert isinstance(target['video_id'], str) and 21 <= int(target['video_id'][1:3]) <= 30
        assert target['target_points_ms'] == []
        for interval in target['candidate_intervals'] + target['context_intervals'] + target['target_intervals']:
            assert isinstance(interval['start_ms'], int) and isinstance(interval['end_ms'], int)
            assert 0 <= interval['start_ms'] < interval['end_ms'] <= target['duration_ms']
            assert interval['evidence_ids'] and set(interval['evidence_ids']) <= ids
            assert interval['boundary_uncertainty_ms'] is None
        assert [(x['start_ms'], x['end_ms']) for x in target['target_intervals']] == [(x['start_ms'], x['end_ms']) for x in target['candidate_intervals']]
        assert all(x['status'] == 'accepted' and x['interval_role'] == 'answer' and x['approval_id'] == approval['approval_id'] for x in target['target_intervals'])
        for point in target['observed_candidate_points']:
            assert 0 <= point['time_ms'] < target['duration_ms'] and point['evidence_id'] in ids
    for source in r['source_records']:
        assert source['sha256'] == sha[source['workbook']]
wb = openpyxl.load_workbook(P / 'ground_truth_review.xlsx', read_only=True, data_only=False)
assert wb.sheetnames == ['Queries', 'Targets', 'Hints', 'Evidence', 'Issues']
assert wb['Queries'].max_row == len(a) + 1
assert wb['Hints'].max_row == sum(len(r['hints']) for r in a) + 1
query_rows = list(wb['Queries'].values)
headers = query_rows[0]
for record, values in zip(a, query_rows[1:]):
    row = dict(zip(headers, values))
    for key in ('query_id', 'annotation_status', 'review_status', 'eligible_for_benchmark'):
        assert row[key] == record[key]
assert wb['Targets'].max_row == 1 + sum(len(t['target_intervals']) + len(t['context_intervals']) for r in a for t in r['targets'])
for sheet in wb:
    assert all(cell.data_type != 'f' for row in sheet for cell in row)
wb.close()
result = {
    'status': 'passed', 'mode': 'metadata_only' if args.metadata_only else 'full_local_artifacts',
    'annotation_records': len(a), 'owner_approved': sum(r['review_status'] == 'approved' for r in a),
    'eligible_tkis_queries': sum(r['eligible_for_benchmark'] for r in a),
    'accepted_target_intervals': sum(len(t['target_intervals']) for r in a for t in r['targets']),
    'source_query_rows_in_inventory': len(inv), 'source_workbook_sha256': sha,
    'local_artifacts_hash_checked': len(checked),
    'checks': ['Unique annotation and source-row IDs', 'Content-bound owner approval',
               'Exact duplicate representative and TKIS-only eligibility', 'Interval bounds and ms units',
               'Evidence references and manifest hashes', 'Exact cumulative hint concatenation',
               'No video IDs or timestamps leaked in hints', 'Workbook/JSONL state agreement',
               'Owner approval does not claim independent continuous-video verification'],
    'not_checked': ['Continuous playback and semantic correctness'] + (['Local artifact bytes', 'Source workbook bytes'] if args.metadata_only else []),
}
if not args.metadata_only:
    (P / 'validation_report.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
print(json.dumps(result, ensure_ascii=False, indent=2))
