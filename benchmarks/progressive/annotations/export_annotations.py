"""Export observed candidate annotations. Never promote sampled images to verified GT."""
from pathlib import Path
import hashlib,json,re,subprocess,sys,platform
from datetime import datetime,timezone
from collections import Counter,defaultdict
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from visual_notes import NOTES
from approval import apply_owner_approval
P=Path(__file__).resolve().parent;ROOT=P.parents[2]
QUEUE=json.loads((P/'review_queue.json').read_text());NOW=datetime.now(timezone.utc).isoformat()
COMMIT=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()
def read_book(name,sheet,required):
 path=ROOT/name;sha=hashlib.sha256(path.read_bytes()).hexdigest();wb=openpyxl.load_workbook(path,read_only=True,data_only=True);headers=None;rows=[]
 for n,row in enumerate(wb[sheet].iter_rows(values_only=True),1):
  if headers is None:
   if required.issubset(set(row)):headers=row
   continue
  vals={str(k):v for k,v in zip(headers,row) if k is not None}
  if any(v is not None for v in vals.values()):rows.append(dict(workbook=name,sha256=sha,sheet=sheet,row=n,values=vals))
 wb.close();assert headers is not None;return rows
base=read_book('TKIS_queries.xlsx','TKIS tổng hợp',{'Query ID','Video ID','Nội dung TKIS'})
frames=read_book('TKIS_queries.xlsx','GT từng frame',{'Query ID','Frame ID','Video ID'})
phm=read_book('TKIS_72_queries_PHM_3_hints.xlsx','72 queries',{'PHM Query ID','Hint 1 (delta)'})
# This ledger lists only pages actually inspected through the vision tool.
seen={n:{'overview_0_0.jpg'} for n in range(1,78)}
for n,j,p in [(2,0,1),(3,1,0),(5,1,0),(9,1,0),(11,1,0),(11,0,1),(16,0,1),(34,0,1),(37,0,1),(46,0,1),(48,0,1),(49,0,1),(50,0,1),(67,0,1),(76,0,1),(77,0,1)]:seen[n].add(f'overview_{j}_{p}.jpg')
extra={22:[('230000_250000',0),('273000_277339',0)],23:[('100000_116000',0)],24:[('673000_711000',0)],25:[('211000_240000',0)],26:[('736000_781000',0)],27:[('158000_265000',0),('158000_265000',1)],32:[('215000_240000',0),('260000_310000',0)],37:[('520000_530000',0)],38:[('590000_740000',0),('590000_740000',1)],39:[('1055000_1145000',0)],42:[('24000_85000',0),('24000_85000',1)],43:[('752000_810000',0)],48:[('192000_240000',0)],50:[('184000_193000',0)],52:[('27000_47000',0)],56:[('1320000_1420000',0),('1320000_1420000',1)],57:[('89000_140000',0)],59:[('610000_725000',0),('610000_725000',1)],60:[('530000_615000',0)],73:[('44000_87000',0)],75:[('440000_505000',0),('440000_505000',1)]}
for n,pages in extra.items():
 for folder,p in pages:seen[n].add(f'extra_{folder}/sheet_{p}.jpg')
(P/'inspection_ledger.json').write_text(json.dumps({str(n):sorted(s) for n,s in seen.items()},ensure_ascii=False,indent=2))
def interval(a,b,duration,ids,kind):
 start=round(a*1000);end=min(round(b*1000),duration);assert 0<=start<end<=duration
 return dict(start_ms=start,end_ms=end,boundary_uncertainty_ms=None,evidence_ids=ids,justification='Khoảng ứng viên từ các ảnh đã quan sát; chưa kiểm tra liên tục và frame sát hai biên.',status='proposed',interval_role=kind)
annotations=[]
for n,q in enumerate(QUEUE,1):
 note=NOTES[n];folder=P/'evidence'/f'{n:03d}';info=json.loads((folder/'inspection.json').read_text());probe=json.loads((folder/'probe.json').read_text());vid=q['video_ids'][0];ev=[]
 for filename in sorted(seen[n]):
  path=folder/filename;assert path.exists(),path
  if filename.startswith('overview_'):
   j,page=map(int,re.search(r'overview_(\d+)_(\d+)',filename).groups());r=info['ranges'][j]
  else:
   r=json.loads((path.parent/'samples.json').read_text());page=int(path.stem.split('_')[-1])
  samples=r['samples'][page*24:(page+1)*24];assert samples
  eid=f'{q["query_id"]}:e{len(ev)+1:02d}'
  ev.append(dict(evidence_id=eid,video_id=vid,start_ms=samples[0]['time_ms'],end_ms=min(info['duration_ms'],samples[-1]['time_ms']+1),modality='visual',observation=note['observation'],observation_scope='Query-level notes; see individual timestamps. Not every statement appears in every frame.',artifact_path=str(path.relative_to(ROOT)),artifact_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),inspection_method='visually_inspected_timestamped_contact_sheet',sample_step_ms=round(r['sample_step_seconds']*1000),sampled_frames=samples,clip_path=r['clip_path'],clip_sha256=r['clip_sha256'],continuous_playback_inspected=False,audio_inspected=False))
 ids=[x['evidence_id'] for x in ev]
 def refs(a,b):
  selected=[e['evidence_id'] for e in ev if e['start_ms']<b*1000 and e['end_ms']>a*1000]
  return selected or ids
 candidates=[interval(a,b,info['duration_ms'],refs(a,b),'answer_candidate') for a,b in note['candidate_intervals_seconds']]
 context=[interval(a,b,info['duration_ms'],refs(a,b),'context_candidate') for a,b in note['context_intervals_seconds']]
 observed=[dict(time_ms=s['time_ms'],evidence_id=e['evidence_id'],artifact_path=s['path']) for e in ev for s in e['sampled_frames'] if any(x['start_ms']<=s['time_ms']<x['end_ms'] for x in candidates)]
 stream=next(s for s in probe['streams'] if s['codec_type']=='video')
 target=dict(video_id=vid,media_url=q['media_urls'][0],media_revision=None,media_revision_status='not_pinned; hashes of decoded evidence retained',duration_ms=info['duration_ms'],source_frame_ids=[str(f) for f in re.findall(r'\d+',str(q['seeds'].get('Frame ID') or ''))],frame_id_kind=info['frame_id_kind'],temporal_mapping=dict(source=info['mapping_path'],source_sha256=info['mapping_sha256'],index_base=None,fps=None,probe=stream,seed_navigation_method=info['seed_mapping'],playback_mapping='reported sample time = seek offset + decoded clip relative PTS; source-frame correspondence and index base not yet verified',source_keyframe_alignment_verified=False),target_points_ms=[],target_intervals=[],context_intervals=context,candidate_intervals=candidates,observed_candidate_points=observed,probe_artifact=str((folder/'probe.json').relative_to(ROOT)))
 hints=[];changes=[]
 for k,delta in enumerate(note['delta_hints'],1):
  old=q['original_hints'].get(f'Hint {k} (delta)')
  hints.append(dict(hint_id=f'{q["query_id"]}:h{k}',delta_text=delta,cumulative_text=' '.join(note['delta_hints'][:k]),evidence_ids=ids,evidence_support_status='provisional_query_level; clause-level review pending',scope='answer_and_context' if context else 'answer_candidate',source='synthetic_visual_sampled',revealed_at_ms=None,review_status='pending'))
  changes.append(dict(hint_number=k,before=old,after=delta,reason='Viết lại theo hình ảnh lấy mẫu; bỏ hoặc sửa chi tiết nguồn chưa được hỗ trợ. '+note['issues'],evidence_ids=ids))
 issues=[dict(code='continuous_video_and_boundary_review_pending',detail='Đã giải mã clip và xem các ảnh lấy mẫu được liệt kê; chưa xem liên tục/kiểm tra sát biên. candidate_intervals không phải ground truth đã chấp nhận.'),dict(code='source_media_alignment_pending',detail='Chưa xác minh revision media và đối chiếu từng keyframe nguồn. Không gán FPS hoặc quy ước index còn chưa chắc.'),dict(code='original_query_clause_review_pending',detail='Hint được rút về chi tiết nhìn thấy; chưa có kiểm chứng đầy đủ từng mệnh đề của query nguồn, nhất là lời nói, danh tính và chuyển động.')]
 if note['issues']:issues.append(dict(code='query_specific_review',detail=note['issues']))
 if n in (28,33):issues.append(dict(code='exact_duplicate',detail='Cùng query, video và seed; synthetic-tkis-028 là đại diện dự kiến, giữ cả hai provenance.'))
 if n==76:issues.append(dict(code='non_tkis_task',detail='Nguồn là TRAKE. Annotation hỗ trợ riêng, chưa chuyển task thành TKIS.'))
 if n==77:issues.append(dict(code='new_media_link_proposed',detail=q['discovery']))
 annotations.append(dict(query_id=q['query_id'],phm_query_id=q['phm_query_id'],source_query_id=q['source_query_id'],source=q['source_records'][0],source_records=q['source_records'],source_match_status=q['source_match_status'],original_query=q['original_query'],query_revision=0,dataset_id='AIC_L21_L30_local_registry',dataset_id_status='local_scope_label_not_official_release_identifier',query_temporal_type='multi_shot_context' if context or len(candidates)>1 else 'single_moment',task_type=q.get('task_type','TKIS'),time_unit='ms',interval_convention='start_inclusive_end_exclusive',targets=[target],evidence=ev,original_hints=q['original_hints'],original_gt_fields=q['seeds'],hint_revision=1,hint_changes=changes,hint_source='synthetic_visual_sampled',hints=hints,annotation_status='needs_review',review_status='pending',annotator='Codex visual sampling',reviewer=None,annotated_at=NOW,code_commit=COMMIT,duplicate_group_id='exact-028-033' if n in (28,33) else None,duplicate_representative_query_id=QUEUE[27]['query_id'] if n in (28,33) else None,connected_video_group_id='video-'+vid,issues=issues,exclusion_reason='Pending continuous media, boundary, source-query and independent review'+('; source task is TRAKE' if n==76 else ''),eligible_for_benchmark=False,review_queue_number=n))
# Each original query lacking a candidate video remains visible in the machine-readable output.
used={(s['workbook'],s['sheet'],s['row']) for a in annotations for s in a['source_records']}
for s in base:
 if (s['workbook'],s['sheet'],s['row']) in used or s['values'].get('Video ID'):continue
 d=s['values'];qid=f'source-{d["Đợt"]}-{d["Query ID"] or "row-"+str(s["row"])}'
 sr=[s]+[f for f in frames if (f['values'].get('Đợt'),f['values'].get('Query ID'))==(d.get('Đợt'),d.get('Query ID'))]
 annotations.append(dict(query_id=qid,source_query_id=d.get('Query ID'),phm_query_id=None,source=s,source_records=sr,source_match_status='confirmed',original_query=d['Nội dung TKIS'],query_revision=0,dataset_id='AIC_L21_L30_local_registry',query_temporal_type=None,task_type='TKIS',time_unit='ms',interval_convention='start_inclusive_end_exclusive',targets=[],evidence=[],original_hints={},hint_source=None,hints=[],annotation_status='unverified',review_status='pending',annotator='Codex source audit',reviewer=None,annotated_at=NOW,duplicate_group_id=None,issues=[dict(code='missing_video_id',detail='Nguồn không có video/frame hợp lệ. Chưa định vị được cảnh nên không tạo start/end hoặc hint giả. Phạm vi L21–L30 chưa xác định.')],exclusion_reason='missing_video_id; unknown_scope',eligible_for_benchmark=False))
# Apply explicit, content-bound owner acceptance; never infer approval from export.
annotations = apply_owner_approval(annotations)
# One inventory record per source query row, not a misleading sum of tasks.
linked=defaultdict(list)
for a in annotations:
 for s in a['source_records']:linked[(s['workbook'],s['sheet'],s['row'])].append(a['query_id'])
frame_lookup=defaultdict(list)
for f in frames:frame_lookup[(f['values'].get('Đợt'),f['values'].get('Query ID'))].append(f)
inventory=[]
for s in base+phm:
 d=s['values'];ids=linked[(s['workbook'],s['sheet'],s['row'])];videos=re.findall(r'[LK]\d+_V\d+',str(d.get('Video ID') or ''))
 if ids:
  statuses={a['annotation_status'] for a in annotations if a['query_id'] in ids};disposition='owner_approved_source_normalized' if 'source_normalized' in statuses else 'annotated_candidates_needs_review' if 'needs_review' in statuses else 'unresolved_video_unknown_scope'
 else:disposition='outside_L21_L30_scope' if videos else 'unresolved_video_unknown_scope'
 rec=dict(source_records=[s],annotation_query_ids=ids,disposition=disposition,source_video_ids=videos,created_at=NOW,code_commit=COMMIT,tool_version='export_annotations_v1',openpyxl_version=openpyxl.__version__)
 if s['workbook']=='TKIS_queries.xlsx':rec['source_records']+=frame_lookup.get((d.get('Đợt'),d.get('Query ID')),[])
 rec['source_conflicts']=[{'frame_sheet_row':f['row'],'code':'query_text_mismatch'} for f in rec['source_records'][1:] if f['values'].get('Nội dung TKIS') and f['values']['Nội dung TKIS']!=d.get('Nội dung TKIS')]
 inventory.append(rec)
# Full frame inventory retains even orphan or conflicting rows.
(P/'source_frames.jsonl').write_text(''.join(json.dumps(f,ensure_ascii=False)+'\n' for f in frames))
(P/'source_inventory.jsonl').write_text(''.join(json.dumps(r,ensure_ascii=False)+'\n' for r in inventory))
(P/'ground_truth.jsonl').write_text(''.join(json.dumps(r,ensure_ascii=False)+'\n' for r in annotations))
# Append-only, content-addressed decision events make repeated exports idempotent.
log=P/'annotation_decisions.jsonl';prior={r['decision_id'] for r in map(json.loads,log.read_text().splitlines())} if log.exists() else set()
with log.open('a') as f:
 for a in annotations:
  digest=hashlib.sha256(json.dumps([a['query_id'],a['targets'],a['hints'],a['issues']],ensure_ascii=False,sort_keys=True).encode()).hexdigest()
  if digest in prior:continue
  event=dict(decision_id=digest,at=NOW,actor=a['annotator'],query_id=a['query_id'],from_status='source_unverified',to_status=a['annotation_status'],changes=['Retain immutable source and source hints','Propose sampled answer intervals separately from accepted intervals','Rewrite three delta hints and regenerate cumulative'] if a['hints'] else ['Retain unresolved source query without invented GT'],reason=a['exclusion_reason'],evidence_ids=[e['evidence_id'] for e in a['evidence']])
  if a.get('approval_id'):
   event.update(actor='dataset_owner',from_status='needs_review',to_status='source_normalized',approval_id=a['approval_id'],changes=['Accept normalized target intervals and synthetic hints for benchmark use','Retain visual sampling provenance','Exclude non-TKIS and non-representative exact duplicates'],reason=a['review_basis'])
  f.write(json.dumps(event,ensure_ascii=False)+'\n')
# Flat workbook with one target/hint/evidence/issue per row.
wb=openpyxl.Workbook();wb.remove(wb.active)
headers={
 'Queries':['query_id','queue_number','source_query_id','phm_query_id','task_type','original_query','annotation_status','review_status','eligible_for_benchmark','duplicate_group_id','source_locations','candidate_interval_count','hint_count','exclusion_reason'],
 'Targets':['query_id','video_id','interval_role','start_ms','end_ms','start_display','end_display','boundary_uncertainty_ms','status','evidence_ids','duration_ms','media_url','justification'],
 'Hints':['query_id','hint_number','delta_text','cumulative_text','original_delta','hint_source','scope','evidence_ids','support_status'],
 'Evidence':['query_id','evidence_id','video_id','start_ms','end_ms','modality','observation','artifact_path','clip_path','sample_step_ms','inspection_method','audio_inspected'],
 'Issues':['query_id','code','detail','source_location']}
for name,hs in headers.items():ws=wb.create_sheet(name);ws.append(hs)
def display(ms):return f'{ms//3600000:02d}:{ms//60000%60:02d}:{ms//1000%60:02d}.{ms%1000:03d}'
for a in annotations:
 loc='; '.join(f'{s["workbook"]} / {s["sheet"]} / row {s["row"]}' for s in a['source_records'] if s['sheet']!='GT từng frame')
 wb['Queries'].append([a['query_id'],a.get('review_queue_number'),a['source_query_id'],a['phm_query_id'],a['task_type'],a['original_query'],a['annotation_status'],a['review_status'],a['eligible_for_benchmark'],a['duplicate_group_id'],loc,sum(len(t['candidate_intervals']) for t in a['targets']),len(a['hints']),a['exclusion_reason']])
 for t in a['targets']:
  for it in (t['target_intervals'] or t['candidate_intervals'])+t['context_intervals']:
   wb['Targets'].append([a['query_id'],t['video_id'],it['interval_role'],it['start_ms'],it['end_ms'],display(it['start_ms']),display(it['end_ms']),it['boundary_uncertainty_ms'],it['status'],'; '.join(it['evidence_ids']),t['duration_ms'],t['media_url'],it['justification']])
 for k,h in enumerate(a['hints'],1):wb['Hints'].append([a['query_id'],k,h['delta_text'],h['cumulative_text'],a['original_hints'].get(f'Hint {k} (delta)'),h['source'],h['scope'],'; '.join(h['evidence_ids']),h['evidence_support_status']])
 for e in a['evidence']:wb['Evidence'].append([a['query_id'],e['evidence_id'],e['video_id'],e['start_ms'],e['end_ms'],e['modality'],e['observation'],e['artifact_path'],e['clip_path'],e['sample_step_ms'],e['inspection_method'],False])
 for issue in a['issues']:wb['Issues'].append([a['query_id'],issue['code'],issue['detail'],loc])
for r in inventory:
 if not r['annotation_query_ids']:
  s=r['source_records'][0];wb['Issues'].append([s['values'].get('Query ID') or s['values'].get('PHM Query ID'),r['disposition'],'Giữ trong inventory; Video ID nguồn ngoài L21–L30.',f'{s["workbook"]} / {s["sheet"]} / row {s["row"]}'])
for ws in wb:
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='243B53')
 for row in ws.iter_rows(min_row=2):
  for c in row:
   c.alignment=Alignment(vertical='top',wrap_text=True)
   if c.data_type=='f':c.data_type='s'
 for col in ws.columns:
  name=col[0].value;ws.column_dimensions[col[0].column_letter].width=65 if name in ('original_query','delta_text','cumulative_text','observation','detail','justification') else 28
 wb.active=0
wb.save(P/'ground_truth_review.xlsx')
counts=dict(source_query_rows=len(base),source_frame_rows=len(frames),phm_query_rows=len(phm),inventory_source_query_records=len(inventory),annotation_records=len(annotations),sampled_queries=len(QUEUE),tkis_sampled_queries=sum(a['task_type']=='TKIS' and bool(a['evidence']) for a in annotations),trake_sampled_queries=1,unique_sampled_videos=len({t['video_id'] for a in annotations for t in a['targets']}),needs_review=sum(a['annotation_status']=='needs_review' for a in annotations),unverified=sum(a['annotation_status']=='unverified' for a in annotations),video_verified=0,approved=sum(a['review_status']=='approved' for a in annotations),eligible_for_benchmark=sum(a['eligible_for_benchmark'] for a in annotations),three_hint_drafts=sum(len(a['hints'])==3 for a in annotations),exact_duplicate_groups=1,exact_duplicate_rows=2,candidate_answer_intervals=sum(len(t['candidate_intervals']) for a in annotations for t in a['targets']),accepted_target_intervals=sum(len(t['target_intervals']) for a in annotations for t in a['targets']),inspected_contact_sheets=sum(len(a['evidence']) for a in annotations),inventory_dispositions=dict(Counter(r['disposition'] for r in inventory)))
(P/'audit_counts.json').write_text(json.dumps(counts,indent=2))
# Reuse local browser editor, preloaded with the autonomous drafts.
page_records=[]
for q,a in zip(QUEUE,annotations):page_records.append({**q,'hints':[h['delta_text'] for h in a['hints']],'candidate_intervals':[{**i,'video_id':t['video_id']} for t in a['targets'] for i in (t['target_intervals'] or t['candidate_intervals'])], 'source_review_status':a['review_status'],'annotation_notes':'\n'.join(i['detail'] for i in a['issues']),'evidence_links':[{'path':str(Path(e['artifact_path']).relative_to(P.relative_to(ROOT))),'label':e['evidence_id']} for e in a['evidence']]})
html=(P/'review_template.html').read_text().replace('__RECORDS__',json.dumps(page_records,ensure_ascii=False).replace('</','<\\/'))
(P/'review.html').write_text(html)
print(json.dumps(counts,ensure_ascii=False,indent=2))
