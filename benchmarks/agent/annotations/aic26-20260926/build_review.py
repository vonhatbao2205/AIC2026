"""Render the frozen review as a local HTML gallery and an Excel label sheet."""
from __future__ import annotations

import html
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from benchmarks.agent.export_submit_dataset import export_review

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]


def local_path(path: str) -> str:
    return str((ROOT / path).relative_to(HERE))


def main():
    rows, audit = export_review(HERE)
    labels = {r["query_id"]: r for r in rows}
    queue = json.loads((HERE / "review_queue.json").read_text())
    decisions = {d["query_id"]: d for d in json.loads((HERE / "review_decisions.json").read_text())["decisions"]}
    evidence = {e["directory"]: e for e in json.loads((HERE / "evidence_index.json").read_text())}
    wb = Workbook()
    ws = wb.active
    ws.title = "Ground truth reviewed"
    ws.append(["Workbook rows", "Query ID", "Task", "Vietnamese query", "Label authority", "Status",
               "Video / time(s) / frame", "Accepted QA answers", "Original DRES verdicts",
               "Visual observation", "Evidence directory", "Split (standalone)"])
    e = html.escape
    parts = ["<!doctype html><html lang='vi'><meta charset='utf-8'><title>AIC26 ground truth review</title>",
             "<style>body{font:16px system-ui;max-width:1500px;margin:30px auto;padding:0 16px;background:#111827;color:#e5e7eb}section{border-top:1px solid #4b5563;padding:20px 0}img{max-width:100%}video{width:460px;max-width:100%}a{color:#93c5fd}pre{white-space:pre-wrap}small{color:#cbd5e1}</style>",
             f"<h1>AIC26 · 26/09/2026 · {audit['queries']} câu đã review</h1>",
             f"<p>{e(str(audit['by_task']))}. Nguồn nhãn: {e(str(audit['by_label_authority']))}.</p>",
             "<p>Ảnh/video bên dưới là bằng chứng local. CORRECT giữ nhãn DRES; WRONG chỉ chọn sau khi đối chiếu theo xác nhận của chủ dữ liệu. Câu đua xe là annotation PE + video, không có verdict DRES. Đây là các mốc điểm, chưa phải toàn bộ khoảng chấp nhận chính thức.</p>"]
    for q in queue:
        d = decisions[q["query_id"]]
        r = labels.get(q["query_id"], {})
        targets = r.get("targets", [])
        moments = "\n".join(f"{t['video_id']} @ {t['start_s']:.6f}s / frame {t.get('frame_idx', '?')}" for t in targets)
        answers = sorted({a for t in targets for a in t.get("answers", [])})
        provenance = r.get("provenance", {})
        ws.append([", ".join(map(str, q["workbook_rows"])), q["query_id"], q["query_type"], q["query"],
                   d["label_authority"], d["review_status"], moments, " | ".join(answers),
                   json.dumps(provenance.get("original_dres_verdicts", {})), d["observation"],
                   d.get("evidence_directory"), r.get("split")])
        parts.extend([f"<section id='row-{q['workbook_rows'][0]}'><h2>Dòng {e(str(q['workbook_rows']))} · {e(q['query_type'])}</h2>",
                      f"<pre>{e(q['query'])}</pre><p>{e(d['label_authority'])} · {e(d['review_status'])}</p>",
                      f"<pre>{e(moments)}</pre><p>QA: {e(' | '.join(answers))}</p><p>{e(d['observation'])}</p>"])
        if d.get("additional_audit_reference"):
            parts.append(f"<p><a href='{e(d['additional_audit_reference'])}'>Audit chi tiết từng hint</a></p>")
        directory = d.get("evidence_directory")
        if directory:
            parts.append(f"<a href='{e(directory)}/overview.jpg'><img loading='lazy' src='{e(directory)}/overview.jpg' alt='Timestamped candidate frames'></a>")
            accepted = set(d.get("accepted_submission_ids", []) + d.get("accepted_annotation_ids", []))
            parts.append("<details><summary>Clip từng candidate (kèm mốc được chọn / loại)</summary>")
            for moment in evidence[directory]["moments"]:
                mid = moment.get("annotation_id") or moment.get("submission_id")
                time = moment.get("annotated_time_s", moment.get("submitted_time_s"))
                clip = local_path(moment["clip_path"])
                parts.append(f"<p>{e(mid)} · E{moment['event_index']} · {time:.6f}s · {'CHỌN' if mid in accepted else 'LOẠI'} · DRES {e(str(moment.get('verdict', 'không có')))}</p><video controls preload='none' src='{e(clip)}'></video>")
            parts.append("</details>")
            for extra in (HERE / directory).glob("extra-*/*.mp4"):
                parts.append(f"<p>Bằng chứng bổ sung: {e(extra.parent.name)}</p><video controls preload='none' src='{e(str(extra.relative_to(HERE)))}'></video>")
        if provenance.get("annotation_details"):
            for annotation in provenance["annotation_details"]:
                parts.append("<details><summary>Tiêu chí ranh giới và sai số annotation</summary>")
                for boundary in annotation["event_boundaries"]:
                    parts.append(f"<p>E{boundary['event_index']}: {e(boundary['criterion'])}<br><small>{e(boundary['uncertainty_note'])}</small></p>")
                    for frame in boundary["evidence_frames"]:
                        parts.append(f"<a href='{e(frame)}'>{e(frame)}</a><br>")
                parts.append("</details>")
        parts.append("</section>")
    parts.append("<h2>Giới hạn</h2><ul>" + "".join(f"<li>{e(s)}</li>" for s in audit["limitations"]) + "</ul></html>")
    (HERE / "review.html").write_text("\n".join(parts))
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1E3A5F")
    for column, width in enumerate([16, 42, 12, 90, 36, 14, 52, 27, 44, 100, 25, 18], 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 120
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(HERE / "ground_truth_review.xlsx")
    print(f"Rendered {len(rows)} reviewed labels")


if __name__ == "__main__":
    main()
